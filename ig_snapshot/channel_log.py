"""Kanal bazlı günlük kayıt: reports/kanallar/<hesap>.xlsx

Her hesap için tek dosya; her snapshot sonrası veritabanından baştan üretilir:
  Günlük           – gün gün takipçi, toplam gönderi, o gün atılan içerik, kazanılan izlenme/beğeni/yorum
  İzlenme (gün)    – gönderi × gün izlenme matrisi (Reels/video)
  Beğeni (gün)     – gönderi × gün beğeni matrisi (tüm gönderiler)
  Gönderiler       – takip edilen tüm gönderilerin son değerleri
"""
from __future__ import annotations

import logging
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from . import config, db
from .content import CONTENT_TYPES
from .report import build_account_report, style_sheet

log = logging.getLogger(__name__)
OUT_DIRNAME = "kanallar"


def _daily_rows(conn, username: str) -> list[list]:
    """Tüm aylar için günlük satırlar (rapor modülündeki hesapla aynı)."""
    daily = {}
    for month in db.months_with_data(conn):
        r = build_account_report(conn, username, month, with_prev=False)
        if r:
            for d in r.daily:
                daily[d.day] = d
    profile = {row["snapshot_date"]: row for row in conn.execute(
        "SELECT snapshot_date, followers_count, media_count FROM profile_snapshots WHERE username = ? ORDER BY 1",
        (username,)).fetchall()}
    tracked = {row["d"]: row["n"] for row in conn.execute("""
        SELECT ms.snapshot_date AS d, COUNT(*) AS n FROM media_snapshots ms
        JOIN media m ON m.media_id = ms.media_id WHERE m.username = ? GROUP BY ms.snapshot_date
    """, (username,)).fetchall()}

    rows = []
    prev_media_count = None
    for day in sorted(set(daily) | set(profile)):
        d = daily.get(day)
        p = profile.get(day)
        media_count = p["media_count"] if p else None
        media_delta = (media_count - prev_media_count
                       if media_count is not None and prev_media_count is not None else None)
        if media_count is not None:
            prev_media_count = media_count
        gr = d.gain_by_group["Reels"] if d else None
        gf = d.gain_by_group["Feed"] if d else None
        rows.append([
            day,
            (p["followers_count"] if p else (d.followers if d else None)),
            d.followers_delta if d else None,
            media_count, media_delta,
            sum(d.posts.values()) if d else 0,
            d.posts_in_group("Reels") if d else 0,
            d.posts_in_group("Feed") if d else 0,
            *[(d.posts.get(t, 0) if d else 0) for t in CONTENT_TYPES if t != "Reels"],
            d.gain.views if d else None, gr.views_or_none if gr else None, gf.views_or_none if gf else None,
            d.gain.likes if d else None, gr.likes if gr else None, gf.likes if gf else None,
            d.gain.comments if d else None,
            tracked.get(day, 0),
        ])
    return rows


DAILY_HEADER = ["Tarih", "Takipçi", "Takipçi Δ", "Toplam gönderi (profil)", "Gönderi Δ",
                "O gün paylaşım", "Reels", "Feed", "Fotoğraf", "Carousel", "Video", "Diğer",
                "Kazanılan izlenme", "Kazanılan izlenme (Reels)", "Kazanılan izlenme (Feed)",
                "Kazanılan beğeni", "Kazanılan beğeni (Reels)", "Kazanılan beğeni (Feed)",
                "Kazanılan yorum", "İzlenen gönderi"]


def _matrix(conn, username: str, since: str) -> tuple[list[str], list, dict, dict]:
    """Son N günün gönderi × gün değerleri. Döner: (günler, gönderiler, izlenme, beğeni)."""
    rows = conn.execute("""
        SELECT ms.media_id, ms.snapshot_date, ms.view_count, ms.like_count
        FROM media_snapshots ms JOIN media m ON m.media_id = ms.media_id
        WHERE m.username = ? AND ms.snapshot_date >= ?
    """, (username, since)).fetchall()
    days = sorted({r["snapshot_date"] for r in rows})
    views: dict[str, dict[str, int]] = defaultdict(dict)
    likes: dict[str, dict[str, int]] = defaultdict(dict)
    for r in rows:
        if r["view_count"] is not None:
            views[r["media_id"]][r["snapshot_date"]] = r["view_count"]
        if r["like_count"] is not None:
            likes[r["media_id"]][r["snapshot_date"]] = r["like_count"]
    ids = set(views) | set(likes)
    media = [m for m in db.media_for_account(conn, username) if m["media_id"] in ids]
    media.sort(key=lambda m: m["published_at"] or "", reverse=True)
    return days, media, views, likes


def _write_matrix(wb, title: str, days: list[str], media, values: dict, only_with_values: bool) -> None:
    ws = wb.create_sheet(title)
    ws.append(["Yayın tarihi", "Tür", "Link", "Açıklama"] + [f"{d[8:10]}.{d[5:7]}" for d in days])
    for m in media:
        vals = values.get(m["media_id"], {})
        if only_with_values and not vals:
            continue
        ws.append([(m["published_at"] or "")[:16].replace("T", " "), m["content_type"], m["permalink"],
                   (m["caption"] or "").replace("\n", " ")[:60]] + [vals.get(d) for d in days])
    style_sheet(ws, [16, 9, 42, 40] + [10] * len(days))
    ws.freeze_panes = "E2"


def write_channel(conn, username: str, path: Path, since: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Günlük"
    ws.append(DAILY_HEADER)
    daily = _daily_rows(conn, username)
    for row in daily:
        ws.append(row)
    style_sheet(ws, [12, 11, 10, 12, 9, 9, 7, 7, 9, 9, 7, 7, 13, 13, 13, 13, 13, 13, 12, 10])
    if len(daily) >= 2:
        chart = LineChart()
        chart.title = "Takipçi"
        chart.height, chart.width = 8, 22
        chart.legend = None
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(chart, f"{get_column_letter(len(DAILY_HEADER) + 2)}2")

    days, media, views, likes = _matrix(conn, username, since)
    _write_matrix(wb, "İzlenme (gün)", days, media, views, only_with_values=True)
    _write_matrix(wb, "Beğeni (gün)", days, media, likes, only_with_values=False)

    ws = wb.create_sheet("Gönderiler")
    ws.append(["Yayın tarihi", "Tür", "İlk görüldü", "Son ölçüm", "İzlenme", "Beğeni", "Yorum", "Link", "Açıklama"])
    latest = {r["media_id"]: r for r in conn.execute("""
        SELECT ms.* FROM media_snapshots ms JOIN media m ON m.media_id = ms.media_id
        WHERE m.username = ? AND ms.snapshot_date = (
            SELECT MAX(x.snapshot_date) FROM media_snapshots x WHERE x.media_id = ms.media_id)
    """, (username,)).fetchall()}
    for m in sorted(db.media_for_account(conn, username), key=lambda m: m["published_at"] or "", reverse=True):
        s = latest.get(m["media_id"])
        ws.append([(m["published_at"] or "")[:16].replace("T", " "), m["content_type"], (m["first_seen"] or "")[:10],
                   s["snapshot_date"] if s else None, s["view_count"] if s else None,
                   s["like_count"] if s else None, s["comments_count"] if s else None,
                   m["permalink"], (m["caption"] or "").replace("\n", " ")[:120]])
    style_sheet(ws, [16, 9, 11, 11, 11, 10, 9, 42, 60])

    wb.save(path)


def generate(conn=None) -> list[Path]:
    own = conn is None
    conn = conn or db.connect()
    try:
        since = (date.today() - timedelta(days=config.CHANNEL_LOG_DAYS)).isoformat()
        out_dir = config.REPORTS_DIR / OUT_DIRNAME
        out_dir.mkdir(parents=True, exist_ok=True)
        written = []
        for a in db.list_accounts(conn):
            u = a["username"]
            if not db.latest_profile(conn, u):
                continue
            path = out_dir / f"{u}.xlsx"
            try:
                write_channel(conn, u, path, since)
                written.append(path)
            except PermissionError:
                log.error("%s yazılamadı — dosya Excel'de açık olabilir", path)
        log.info("Kanal dosyaları yazıldı: %d (%s)", len(written), out_dir)
        return written
    finally:
        if own:
            conn.close()
