"""Günlük içerik raporu: reports/YYYY-MM/gunluk-YYYY-MM.xlsx

Kesit = "o gün yayınlanan içerik" (kohort). Her ay sıfırdan başlar, her gece yeniden üretilir.
  Günlük içerik      – hesap × gün: kaç içerik atıldı, o içerikler bugüne kadar ne izlenme/beğeni topladı,
                       ilk gün (yayın günü akşamı) kaçtaydı, en iyi gönderi
  Reels günlük /
  Feed günlük        – aynı tablo grup bazında
  Pivot paylaşım /
  Pivot izlenme /
  Pivot beğeni       – gün × hesap karşılaştırma (grafikli)
  Gönderiler         – ay içindeki her gönderi: ilk gün / ay sonu / güncel (sürüklenme dahil) / son 24 saat artışı
  Öne çıkanlar       – güncel izlenme, son 24 saat artışı ve beğeniye göre ilk 20
"""
from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from . import config, db
from .content import CONTENT_TYPES, GROUPS, group_of
from .report import month_bounds, month_label, ordered_usernames, style_sheet

log = logging.getLogger(__name__)
TOP_N = 20


@dataclass
class PostPerf:
    media_id: str
    username: str
    published_at: str          # yerel ISO
    content_type: str
    caption: str
    permalink: str
    first_date: str | None = None
    first_views: int | None = None
    first_likes: int | None = None
    first_comments: int | None = None
    last_date: str | None = None
    last_views: int | None = None
    last_likes: int | None = None
    last_comments: int | None = None
    prev_views: int | None = None
    prev_likes: int | None = None
    end_date: str | None = None      # ay sonuna kadarki son ölçüm (sürüklenme öncesi)
    end_views: int | None = None
    end_likes: int | None = None
    end_comments: int | None = None

    @property
    def drift_views(self) -> int | None:
        """Ay bittikten sonra gelen izlenme (güncel − ay sonu)."""
        if self.last_views is None or self.end_views is None:
            return None
        return self.last_views - self.end_views

    @property
    def day(self) -> str:
        return self.published_at[:10]

    @property
    def group(self) -> str:
        return group_of(self.content_type)

    @property
    def gain24_views(self) -> int | None:
        if self.last_views is None or self.prev_views is None:
            return None
        return self.last_views - self.prev_views

    @property
    def gain24_likes(self) -> int | None:
        if self.last_likes is None or self.prev_likes is None:
            return None
        return self.last_likes - self.prev_likes


@dataclass
class Cohort:
    """Bir hesabın belirli bir günde yayınladığı içerikler."""
    posts: list[PostPerf] = field(default_factory=list)

    def count(self, t: str | None = None, group: str | None = None) -> int:
        return sum(1 for p in self.posts
                   if (t is None or p.content_type == t) and (group is None or p.group == group))

    def _sum(self, attr: str, group: str | None = None) -> int | None:
        vals = [getattr(p, attr) for p in self.posts if group is None or p.group == group]
        vals = [v for v in vals if v is not None]
        return sum(vals) if vals else None

    def _avg(self, attr: str, group: str | None = None) -> float | None:
        vals = [getattr(p, attr) for p in self.posts if group is None or p.group == group]
        vals = [v for v in vals if v is not None]
        return sum(vals) / len(vals) if vals else None

    def best(self, group: str | None = None) -> PostPerf | None:
        cands = [p for p in self.posts if group is None or p.group == group]
        if not cands:
            return None
        return max(cands, key=lambda p: (p.last_views or 0, p.last_likes or 0))


def load_posts(conn, username: str, month: str) -> list[PostPerf]:
    rows = conn.execute("""
        SELECT media_id, published_at, content_type, caption, permalink FROM media
        WHERE username = ? AND published_month = ? AND published_at IS NOT NULL
    """, (username, month)).fetchall()
    posts = {r["media_id"]: PostPerf(r["media_id"], username, r["published_at"], r["content_type"] or "Diğer",
                                     (r["caption"] or "").replace("\n", " ").strip(), r["permalink"] or "")
             for r in rows}
    if not posts:
        return []
    snaps = conn.execute(f"""
        SELECT media_id, snapshot_date, view_count, like_count, comments_count FROM media_snapshots
        WHERE media_id IN ({",".join("?" * len(posts))}) ORDER BY media_id, snapshot_date
    """, list(posts)).fetchall()
    _, month_end = month_bounds(month)
    for s in snaps:
        p = posts[s["media_id"]]
        if p.first_date is None:
            p.first_date, p.first_views, p.first_likes, p.first_comments = (
                s["snapshot_date"], s["view_count"], s["like_count"], s["comments_count"])
        p.prev_views, p.prev_likes = p.last_views, p.last_likes
        p.last_date, p.last_views, p.last_likes, p.last_comments = (
            s["snapshot_date"], s["view_count"], s["like_count"], s["comments_count"])
        if s["snapshot_date"] <= month_end:
            p.end_date, p.end_views, p.end_likes, p.end_comments = (
                s["snapshot_date"], s["view_count"], s["like_count"], s["comments_count"])
    # tek ölçümü olan gönderide "önceki" = ilk gün öncesi (0) — 24 saat artışı = tamamı
    for p in posts.values():
        if p.first_date == p.last_date:
            p.prev_views = 0 if p.last_views is not None else None
            p.prev_likes = 0 if p.last_likes is not None else None
    return sorted(posts.values(), key=lambda p: p.published_at, reverse=True)


def month_days(month: str) -> list[str]:
    start, end = month_bounds(month)
    today = date.today().isoformat()
    end = min(end, today)
    d = date.fromisoformat(start)
    out = []
    while d.isoformat() <= end:
        out.append(d.isoformat())
        d = d.fromordinal(d.toordinal() + 1)
    return out


def _r(x) -> int | None:
    return round(x) if x is not None else None


def _cohort_row(day: str, username: str, c: Cohort, group: str | None) -> list:
    best = c.best(group)
    reels_avg = c._avg("last_views", "Reels") if group in (None, "Reels") else None
    return [day, username, c.count(group=group),
            *([c.count(t) for t in CONTENT_TYPES if t != "Diğer"] if group is None else []),
            c._sum("last_views", group), c._sum("first_views", group),
            c._sum("last_likes", group), c._sum("first_likes", group),
            c._sum("last_comments", group),
            _r(reels_avg), _r(c._avg("last_likes", group)),
            best.last_views if best else None, best.last_likes if best else None,
            best.permalink if best else None, (best.caption[:60] if best else None)]


def _cohort_header(group: str | None) -> list[str]:
    types = [t for t in CONTENT_TYPES if t != "Diğer"] if group is None else []
    return (["Tarih", "Hesap", "Paylaşım"] + types +
            ["İzlenme (güncel)", "İzlenme (ilk gün)", "Beğeni (güncel)", "Beğeni (ilk gün)", "Yorum",
             "Ort. izlenme / Reels", "Ort. beğeni / gönderi",
             "En iyi: izlenme", "En iyi: beğeni", "En iyi: link", "En iyi: açıklama"])


def _pivot(wb, title: str, days: list[str], usernames: list[str], values: dict, chart_title: str) -> None:
    ws = wb.create_sheet(title)
    ws.append(["Tarih"] + usernames)
    for d in days:
        ws.append([d] + [values.get((d, u)) for u in usernames])
    style_sheet(ws, [12] + [15] * len(usernames))
    if len(days) >= 2:
        ch = LineChart()
        ch.title = chart_title
        ch.height, ch.width = 10, 24
        ch.legend.position = "b"
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_col=1 + len(usernames), max_row=ws.max_row),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(ch, f"{get_column_letter(len(usernames) + 3)}2")


def write_xlsx(month: str, cohorts: dict[str, dict[str, Cohort]], posts: list[PostPerf], path: Path) -> None:
    usernames = list(cohorts)
    days = month_days(month)
    wb = Workbook()

    for title, group in (("Günlük içerik", None), ("Reels günlük", "Reels"), ("Feed günlük", "Feed")):
        ws = wb.active if group is None else wb.create_sheet(title)
        ws.title = title
        ws.append(_cohort_header(group))
        for d in days:
            for u in usernames:
                ws.append(_cohort_row(d, u, cohorts[u].get(d, Cohort()), group))
        widths = [12, 18, 9] + ([7] * (len(CONTENT_TYPES) - 1) if group is None else []) + \
                 [14, 14, 13, 13, 9, 14, 14, 13, 12, 42, 40]
        style_sheet(ws, widths)

    piv_posts, piv_views, piv_likes = {}, {}, {}
    for u in usernames:
        for d in days:
            c = cohorts[u].get(d)
            piv_posts[(d, u)] = c.count() if c else 0
            piv_views[(d, u)] = c._sum("last_views") if c else None
            piv_likes[(d, u)] = c._sum("last_likes") if c else None
    _pivot(wb, "Pivot paylaşım", days, usernames, piv_posts, "Günlük paylaşım sayısı")
    _pivot(wb, "Pivot izlenme", days, usernames, piv_views, "O gün atılan içeriklerin izlenmesi (güncel)")
    _pivot(wb, "Pivot beğeni", days, usernames, piv_likes, "O gün atılan içeriklerin beğenisi (güncel)")

    ws = wb.create_sheet("Gönderiler")
    ws.append(["Yayın", "Hesap", "Grup", "Tür", "İzlenme (ilk gün)", "İzlenme (ay sonu)", "İzlenme (güncel)",
               "Ay sonrası sürüklenme", "Son 24s izlenme Δ",
               "Beğeni (ilk gün)", "Beğeni (ay sonu)", "Beğeni (güncel)", "Son 24s beğeni Δ", "Yorum",
               "İlk ölçüm", "Son ölçüm", "Açıklama", "Link"])
    for p in posts:
        ws.append([p.published_at[:16].replace("T", " "), p.username, p.group, p.content_type,
                   p.first_views, p.end_views, p.last_views, p.drift_views, p.gain24_views,
                   p.first_likes, p.end_likes, p.last_likes, p.gain24_likes,
                   p.last_comments, p.first_date, p.last_date, p.caption[:120], p.permalink])
    style_sheet(ws, [16, 16, 7, 9, 13, 13, 13, 13, 14, 12, 12, 12, 14, 8, 11, 11, 60, 42])

    ws = wb.create_sheet("Öne çıkanlar")
    blocks = [
        ("En çok izlenen (güncel)", sorted([p for p in posts if p.last_views], key=lambda p: -p.last_views)[:TOP_N],
         lambda p: p.last_views),
        ("Son 24 saatte en çok izlenme kazanan", sorted([p for p in posts if p.gain24_views], key=lambda p: -p.gain24_views)[:TOP_N],
         lambda p: p.gain24_views),
        ("En çok beğenilen (güncel)", sorted([p for p in posts if p.last_likes], key=lambda p: -p.last_likes)[:TOP_N],
         lambda p: p.last_likes),
    ]
    ws.append(["Liste", "#", "Yayın", "Hesap", "Tür", "Değer", "İzlenme", "Beğeni", "Yorum", "Açıklama", "Link"])
    for title, items, key in blocks:
        for i, p in enumerate(items, 1):
            ws.append([title, i, p.published_at[:16].replace("T", " "), p.username, p.content_type, key(p),
                       p.last_views, p.last_likes, p.last_comments, p.caption[:100], p.permalink])
    style_sheet(ws, [34, 4, 16, 16, 9, 12, 12, 10, 8, 60, 42])

    wb.save(path)


def build(conn, month: str) -> tuple[dict[str, dict[str, Cohort]], list[PostPerf]]:
    start, end = month_bounds(month)
    cohorts: dict[str, dict[str, Cohort]] = {}
    all_posts: list[PostPerf] = []
    for u in ordered_usernames(conn, start, end):
        posts = load_posts(conn, u, month)
        by_day: dict[str, Cohort] = defaultdict(Cohort)
        for p in posts:
            by_day[p.day].posts.append(p)
        cohorts[u] = dict(by_day)
        all_posts.extend(posts)
    all_posts.sort(key=lambda p: p.published_at, reverse=True)
    return cohorts, all_posts


def generate(month: str, conn=None) -> Path | None:
    own = conn is None
    conn = conn or db.connect()
    try:
        cohorts, posts = build(conn, month)
    finally:
        if own:
            conn.close()
    if not cohorts:
        return None
    out_dir = config.REPORTS_DIR / month
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"gunluk-{month}.xlsx"
    try:
        write_xlsx(month, cohorts, posts, path)
    except PermissionError:
        log.error("%s yazılamadı — dosya Excel'de açık olabilir", path)
        return None
    log.info("Günlük içerik raporu yazıldı: %s (%d gönderi)", path, len(posts))
    return path


def highlights_today(conn, day: date, n: int = 5) -> dict[str, list[PostPerf]]:
    """Bugün yayınlanan içerikler: en çok izlenen n Reels ve en çok beğenilen n Feed (Telegram için)."""
    month = day.strftime("%Y-%m")
    _, posts = build(conn, month)
    today = day.isoformat()
    todays = [p for p in posts if p.day == today]
    reels = [p for p in todays if p.group == "Reels"]
    feed = [p for p in todays if p.group == "Feed"]
    return {
        "reels": sorted(reels, key=lambda p: (p.last_views or 0, p.last_likes or 0), reverse=True)[:n],
        "feed": sorted(feed, key=lambda p: (p.last_likes or 0, p.last_comments or 0), reverse=True)[:n],
    }
