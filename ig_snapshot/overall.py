"""Genel görünüm: takip başından bugüne ay ay takipçi, Reels/Feed paylaşım, beğeni ve izlenme artışı."""
from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

from . import config, db
from .content import GROUPS
from .report import (AccountReport, Gain, build_account_report, daily_header, daily_row, daily_widths,
                     fmt, fmt_day, fmt_delta, fmt_pct, month_label, ordered_usernames, style_sheet, _round)

log = logging.getLogger(__name__)
OUT_DIRNAME = "genel"


@dataclass
class AccountHistory:
    username: str
    name: str | None
    months: list[AccountReport]          # kronolojik
    first_date: str
    first_followers: int | None
    last_date: str
    last_followers: int | None

    @property
    def total_delta(self) -> int | None:
        if self.first_followers is None or self.last_followers is None:
            return None
        return self.last_followers - self.first_followers

    @property
    def total_pct(self) -> float | None:
        if self.total_delta is None or not self.first_followers:
            return None
        return self.total_delta / self.first_followers * 100

    def sum_posts(self, group: str | None = None) -> int:
        return sum(m.group_count(group) if group else m.total_posts for m in self.months)

    def sum_gain(self, group: str | None = None) -> Gain:
        total = Gain()
        for m in self.months:
            total.add(m.gain_by_group[group] if group else m.gain)
        return total

    def avg_likes(self, group: str) -> float | None:
        likes = sum(m.by_group[group].likes for m in self.months)
        known = sum(m.by_group[group].likes_known for m in self.months)
        return likes / known if known else None


def build_overall(conn) -> list[AccountHistory]:
    months = db.months_with_data(conn)
    if not months:
        return []
    usernames = ordered_usernames(conn, f"{months[0]}-01", f"{months[-1]}-31")
    histories: list[AccountHistory] = []
    for u in usernames:
        reports = [r for r in (build_account_report(conn, u, m, with_prev=False) for m in months) if r]
        if not reports:
            continue
        first = conn.execute(
            "SELECT snapshot_date, followers_count FROM profile_snapshots WHERE username = ? "
            "ORDER BY snapshot_date LIMIT 1", (u,)).fetchone()
        last = db.latest_profile(conn, u)
        histories.append(AccountHistory(
            username=u, name=reports[-1].name, months=reports,
            first_date=first["snapshot_date"], first_followers=first["followers_count"],
            last_date=last["snapshot_date"], last_followers=last["followers_count"],
        ))
    return histories


# --- Markdown -------------------------------------------------------------------

def _split(total: str, reels: str, feed: str) -> str:
    return f"{total} ({reels} / {feed})"


def write_markdown(histories: list[AccountHistory], path: Path) -> None:
    now = datetime.now()
    first = min(h.first_date for h in histories)
    last = max(h.last_date for h in histories)
    n_months = len({m.month for h in histories for m in h.months})
    lines = [
        "# Instagram Rakip Takibi — Genel Görünüm",
        "",
        f"Oluşturulma: {now:%d.%m.%Y %H:%M} · Takip başlangıcı: {fmt_day(first)} · Son ölçüm: {fmt_day(last)} · "
        f"{n_months} ay · {len(histories)} hesap",
        "",
        "> Parantez içleri **(Reels / Feed)** kırılımı. Kazanılan değerler, takip edilen tüm gönderilerin "
        "o dönemdeki artışıdır. Ayrıntılı Excel: `genel-rapor.xlsx` (takipçi artış grafiği, aylık izlenme/beğeni grafikleri).",
        "",
        "## Takip başından bugüne",
        "",
        "| Hesap | İlk ölçüm | Takipçi (ilk) | Takipçi (son) | Toplam artış | Paylaşım (Reels / Feed) | "
        "Kazanılan izlenme (Reels / Feed) | Kazanılan beğeni (Reels / Feed) | Ort. beğeni Reels | Ort. beğeni Feed |",
        "|---|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for h in histories:
        gr, gf, g = h.sum_gain("Reels"), h.sum_gain("Feed"), h.sum_gain()
        lines.append(
            f"| @{h.username} | {fmt_day(h.first_date)} | {fmt(h.first_followers)} | {fmt(h.last_followers)} | "
            f"{fmt_delta(h.total_delta)} ({fmt_pct(h.total_pct)}) | "
            f"{_split(str(h.sum_posts()), str(h.sum_posts('Reels')), str(h.sum_posts('Feed')))} | "
            f"{_split(fmt_delta(g.views), fmt_delta(gr.views_or_none), fmt_delta(gf.views_or_none))} | "
            f"{_split(fmt_delta(g.likes), fmt_delta(gr.likes), fmt_delta(gf.likes))} | "
            f"{fmt(h.avg_likes('Reels'))} | {fmt(h.avg_likes('Feed'))} |"
        )

    for h in histories:
        title = f"@{h.username}" + (f" — {h.name}" if h.name else "")
        lines += ["", f"## {title}", "",
                  "| Ay | Takipçi (ay sonu) | Takipçi Δ | Paylaşım (Reels / Feed) | Kazanılan izlenme (Reels / Feed) | "
                  "Kazanılan beğeni (Reels / Feed) | Kazanılan yorum | Ort. beğeni Reels | Ort. beğeni Feed | Ölçüm günü |",
                  "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|"]
        for m in h.months:
            gr, gf = m.gain_by_group["Reels"], m.gain_by_group["Feed"]
            lines.append(
                f"| {month_label(m.month)} | {fmt(m.followers_end)} | {fmt_delta(m.followers_delta)} ({fmt_pct(m.followers_pct)}) | "
                f"{_split(str(m.total_posts), str(m.group_count('Reels')), str(m.group_count('Feed')))} | "
                f"{_split(fmt_delta(m.gain.views), fmt_delta(gr.views_or_none), fmt_delta(gf.views_or_none))} | "
                f"{_split(fmt_delta(m.gain.likes), fmt_delta(gr.likes), fmt_delta(gf.likes))} | "
                f"{fmt_delta(m.gain.comments)} | {fmt(m.by_group['Reels'].avg_likes)} | "
                f"{fmt(m.by_group['Feed'].avg_likes)} | {m.snapshot_days} |"
            )
        gr, gf, g = h.sum_gain("Reels"), h.sum_gain("Feed"), h.sum_gain()
        lines.append(
            f"| **Toplam** | {fmt(h.last_followers)} | **{fmt_delta(h.total_delta)}** ({fmt_pct(h.total_pct)}) | "
            f"{_split(str(h.sum_posts()), str(h.sum_posts('Reels')), str(h.sum_posts('Feed')))} | "
            f"{_split(fmt_delta(g.views), fmt_delta(gr.views_or_none), fmt_delta(gf.views_or_none))} | "
            f"{_split(fmt_delta(g.likes), fmt_delta(gr.likes), fmt_delta(gf.likes))} | "
            f"{fmt_delta(g.comments)} | {fmt(h.avg_likes('Reels'))} | {fmt(h.avg_likes('Feed'))} | "
            f"{sum(m.snapshot_days for m in h.months)} |"
        )

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# --- Excel ------------------------------------------------------------------------

def _pivot_sheet(wb, title: str, row_labels: list[str], col_labels: list[str], values: dict,
                 chart_title: str, chart_cls, y_title: str, row_header: str) -> None:
    """Satır × sütun tablosu + yanına grafik. values[(row, col)] -> sayı."""
    ws = wb.create_sheet(title)
    ws.append([row_header] + col_labels)
    for rl in row_labels:
        ws.append([rl] + [values.get((rl, cl)) for cl in col_labels])
    style_sheet(ws, [14] + [16] * len(col_labels))
    if len(row_labels) < 2:
        return
    chart = chart_cls()
    chart.title = chart_title
    chart.y_axis.title = y_title
    chart.height, chart.width = 11, 26
    chart.legend.position = "b"
    data = Reference(ws, min_col=2, min_row=1, max_col=1 + len(col_labels), max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    if chart_cls is LineChart:
        for s in chart.series:
            s.smooth = False
            s.marker.symbol = "none"
    ws.add_chart(chart, f"{get_column_letter(len(col_labels) + 3)}2")


def write_xlsx(conn, histories: list[AccountHistory], path: Path) -> None:
    wb = Workbook()
    usernames = [h.username for h in histories]

    # Aylık tablo
    ws = wb.active
    ws.title = "Aylık"
    ws.append(["Hesap", "Ay", "Takipçi başlangıç", "Takipçi son", "Takipçi Δ", "Takipçi %",
               "Paylaşım", "Reels paylaşım", "Feed paylaşım",
               "Kazanılan izlenme", "Kazanılan izlenme (Reels)", "Kazanılan izlenme (Feed)",
               "Kazanılan beğeni", "Kazanılan beğeni (Reels)", "Kazanılan beğeni (Feed)", "Kazanılan yorum",
               "Ort. beğeni Reels", "Ort. beğeni Feed", "Ort. izlenme Reels", "Ort. izlenme Feed", "Ölçüm günü"])
    for h in histories:
        for m in h.months:
            gr, gf = m.gain_by_group["Reels"], m.gain_by_group["Feed"]
            br, bf = m.by_group["Reels"], m.by_group["Feed"]
            ws.append([h.username, m.month, m.followers_start, m.followers_end, m.followers_delta,
                       round(m.followers_pct, 2) if m.followers_pct is not None else None,
                       m.total_posts, br.count, bf.count,
                       m.gain.views, gr.views_or_none, gf.views_or_none, m.gain.likes, gr.likes, gf.likes, m.gain.comments,
                       _round(br.avg_likes), _round(bf.avg_likes), _round(br.avg_views), _round(bf.avg_views),
                       m.snapshot_days])
    style_sheet(ws, [18, 9, 13, 13, 11, 10, 10, 10, 10, 14, 14, 14, 14, 14, 14, 13, 12, 12, 12, 12, 9])

    # Takipçi trendi (mutlak) ve artış (ilk ölçüme göre) — günlük
    rows = conn.execute(
        "SELECT snapshot_date, username, followers_count FROM profile_snapshots ORDER BY snapshot_date"
    ).fetchall()
    dates = sorted({r["snapshot_date"] for r in rows})
    absolute: dict = {}
    growth: dict = {}
    first_val = {h.username: h.first_followers for h in histories}
    for r in rows:
        u = r["username"]
        if u not in first_val or r["followers_count"] is None:
            continue
        absolute[(r["snapshot_date"], u)] = r["followers_count"]
        if first_val[u] is not None:
            growth[(r["snapshot_date"], u)] = r["followers_count"] - first_val[u]
    _pivot_sheet(wb, "Takipçi artışı", dates, usernames, growth,
                 "Takipçi artışı (ilk ölçüme göre)", LineChart, "Takipçi Δ", "Tarih")
    _pivot_sheet(wb, "Takipçi", dates, usernames, absolute,
                 "Takipçi sayısı", LineChart, "Takipçi", "Tarih")

    # Aylık kazanılan izlenme / beğeni (hesap bazında sütun grafik)
    months = sorted({m.month for h in histories for m in h.months})
    month_rows = [month_label(m) for m in months]
    views_v: dict = {}
    likes_v: dict = {}
    for h in histories:
        for m in h.months:
            views_v[(month_label(m.month), h.username)] = m.gain.views
            likes_v[(month_label(m.month), h.username)] = m.gain.likes
    _pivot_sheet(wb, "Aylık izlenme", month_rows, usernames, views_v,
                 "Aylık kazanılan izlenme", BarChart, "İzlenme", "Ay")
    _pivot_sheet(wb, "Aylık beğeni", month_rows, usernames, likes_v,
                 "Aylık kazanılan beğeni", BarChart, "Beğeni", "Ay")

    # Reels / Feed aylık kırılım (hesap × ay)
    ws = wb.create_sheet("Reels-Feed aylık")
    ws.append(["Hesap", "Ay", "Grup", "Paylaşım", "Toplam izlenme", "Ort. izlenme", "Toplam beğeni", "Ort. beğeni",
               "Kazanılan izlenme", "Kazanılan beğeni", "Kazanılan yorum"])
    for h in histories:
        for m in h.months:
            for g in GROUPS:
                ts, gn = m.by_group[g], m.gain_by_group[g]
                ws.append([h.username, m.month, g, ts.count, ts.views_or_none, _round(ts.avg_views), ts.likes,
                           _round(ts.avg_likes), gn.views_or_none, gn.likes, gn.comments])
    style_sheet(ws, [18, 9, 8, 10, 14, 12, 13, 11, 15, 15, 15])

    # Tüm zamanlar günlük
    ws = wb.create_sheet("Günlük (tümü)")
    ws.append(daily_header())
    for h in histories:
        for m in h.months:
            for d in m.daily:
                ws.append(daily_row(h.username, d))
    style_sheet(ws, daily_widths())

    wb.save(path)


def generate(conn=None) -> dict | None:
    own = conn is None
    conn = conn or db.connect()
    try:
        histories = build_overall(conn)
        if not histories:
            log.info("Genel rapor için veri yok")
            return None
        out_dir = config.REPORTS_DIR / OUT_DIRNAME
        out_dir.mkdir(parents=True, exist_ok=True)
        md = out_dir / "genel-rapor.md"
        xlsx = out_dir / "genel-rapor.xlsx"
        write_markdown(histories, md)
        write_xlsx(conn, histories, xlsx)
    finally:
        if own:
            conn.close()
    log.info("Genel rapor yazıldı: %s", out_dir)
    return {"dir": out_dir, "md": md, "xlsx": xlsx, "histories": histories}
