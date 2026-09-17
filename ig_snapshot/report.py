"""Aylık rapor: takipçi artışı, Reels/Feed ve tür bazında içerik üretimi, beğeni ve izlenmeler."""
from __future__ import annotations

import calendar
import csv
import logging
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db
from .content import CONTENT_TYPES, GROUPS, group_of

log = logging.getLogger(__name__)

MONTHS_TR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
             "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
TOP_N = 5


# --- yardımcılar --------------------------------------------------------------

def month_bounds(month: str) -> tuple[str, str]:
    y, m = (int(x) for x in month.split("-"))
    return f"{y:04d}-{m:02d}-01", f"{y:04d}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}"


def prev_month(month: str) -> str:
    y, m = (int(x) for x in month.split("-"))
    return f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"


def month_label(month: str) -> str:
    y, m = (int(x) for x in month.split("-"))
    return f"{MONTHS_TR[m - 1]} {y}"


def fmt(n) -> str:
    if n is None:
        return "–"
    if isinstance(n, float):
        return f"{n:,.0f}".replace(",", ".")
    return f"{n:,}".replace(",", ".")


def fmt_delta(n) -> str:
    if n is None:
        return "–"
    return ("+" if n > 0 else "") + fmt(n)


def fmt_pct(x) -> str:
    if x is None:
        return "–"
    return ("+" if x > 0 else "") + f"%{x:.1f}".replace(".", ",")


def fmt_day(iso: str | None) -> str:
    if not iso:
        return "–"
    return f"{iso[8:10]}.{iso[5:7]}.{iso[0:4]}"


def _avg(total: int, known: int) -> float | None:
    return total / known if known else None


def _round(x: float | None) -> int | None:
    return round(x) if x is not None else None


# --- veri modelleri -----------------------------------------------------------

@dataclass
class Gain:
    """Bir dönemde kazanılan izlenme / beğeni / yorum."""
    views: int = 0
    likes: int = 0
    comments: int = 0
    views_seen: bool = False   # bu grupta en az bir gönderinin izlenmesi ölçülebildi mi

    def add(self, other: "Gain") -> None:
        self.views += other.views
        self.likes += other.likes
        self.comments += other.comments
        self.views_seen = self.views_seen or other.views_seen

    @property
    def views_or_none(self) -> int | None:
        return self.views if self.views_seen else None


@dataclass
class TypeStats:
    """Belirli bir tür/grupta bu ay yayınlanan gönderilerin son ölçülen toplamları."""
    count: int = 0
    views: int = 0
    views_known: int = 0
    likes: int = 0
    likes_known: int = 0
    comments: int = 0
    comments_known: int = 0

    @property
    def avg_views(self) -> float | None:
        return _avg(self.views, self.views_known)

    @property
    def avg_likes(self) -> float | None:
        return _avg(self.likes, self.likes_known)

    @property
    def avg_comments(self) -> float | None:
        return _avg(self.comments, self.comments_known)

    @property
    def views_or_none(self) -> int | None:
        """İzlenme hiç ölçülemediyse (fotoğraf/carousel) 0 yerine None."""
        return self.views if self.views_known else None

    def add_post(self, views: int | None, likes: int | None, comments: int | None) -> None:
        self.count += 1
        if views is not None:
            self.views += views
            self.views_known += 1
        if likes is not None:
            self.likes += likes
            self.likes_known += 1
        if comments is not None:
            self.comments += comments
            self.comments_known += 1


@dataclass
class PostRow:
    media_id: str
    content_type: str
    published_at: str
    views: int | None
    likes: int | None
    comments: int | None
    caption: str
    permalink: str

    @property
    def group(self) -> str:
        return group_of(self.content_type)


@dataclass
class DailyRow:
    day: str
    followers: int | None
    followers_delta: int | None
    posts: dict[str, int]                 # türe göre o gün yayınlanan gönderi
    gain: Gain                            # o gün kazanılan toplam
    gain_by_group: dict[str, Gain]        # Reels / Feed

    def posts_in_group(self, group: str) -> int:
        return sum(n for t, n in self.posts.items() if group_of(t) == group)


@dataclass
class AccountReport:
    username: str
    name: str | None
    month: str
    followers_start: int | None
    followers_start_date: str | None
    followers_end: int | None
    followers_end_date: str | None
    by_type: dict[str, TypeStats]
    by_group: dict[str, TypeStats]
    posts: list[PostRow]
    gain: Gain
    gain_by_group: dict[str, Gain]
    daily: list[DailyRow]
    snapshot_days: int
    prev: dict | None = None
    view_count_available: bool = True

    @property
    def followers_delta(self) -> int | None:
        if self.followers_start is None or self.followers_end is None:
            return None
        return self.followers_end - self.followers_start

    @property
    def followers_pct(self) -> float | None:
        if self.followers_delta is None or not self.followers_start:
            return None
        return self.followers_delta / self.followers_start * 100

    @property
    def total_posts(self) -> int:
        return sum(t.count for t in self.by_type.values())

    @property
    def total_views(self) -> int:
        return sum(t.views for t in self.by_type.values())

    @property
    def avg_views(self) -> float | None:
        return _avg(self.total_views, sum(t.views_known for t in self.by_type.values()))

    def type_count(self, t: str) -> int:
        return self.by_type[t].count if t in self.by_type else 0

    def group_count(self, g: str) -> int:
        return self.by_group[g].count

    @property
    def top_posts(self) -> list[PostRow]:
        return sorted(self.posts, key=lambda p: (p.views or 0, p.likes or 0), reverse=True)[:TOP_N]


# --- hesaplama ----------------------------------------------------------------

def build_account_report(conn, username: str, month: str, with_prev: bool = True) -> AccountReport | None:
    start, end = month_bounds(month)
    series = db.profile_series(conn, username, start, end)
    if not series:
        return None
    info = db.account_info(conn, username)
    baseline_profile = db.last_profile_before(conn, username, start)
    start_row = baseline_profile or series[0]
    end_row = series[-1]

    media_rows = db.media_for_account(conn, username)
    media_by_id = {r["media_id"]: r for r in media_rows}
    snaps = db.media_snapshots_in_range(conn, username, start, end)
    baseline = db.media_baseline_before(conn, username, start)

    # Gün gün kazanılan izlenme/beğeni/yorum (Reels ve Feed ayrı):
    # her gönderinin bir önceki ölçüme göre artışı
    latest: dict[str, object] = {}
    last_val: dict[str, tuple] = {}
    daily_gain: dict[str, dict[str, Gain]] = defaultdict(lambda: {g: Gain() for g in GROUPS})
    any_view = False
    for s in snaps:
        mid = s["media_id"]
        mrow = media_by_id.get(mid)
        cur = (s["view_count"], s["like_count"], s["comments_count"])
        if cur[0] is not None:
            any_view = True
        if mid not in last_val:
            if mid in baseline:
                b = baseline[mid]
                last_val[mid] = (b["view_count"], b["like_count"], b["comments_count"])
            elif mrow is not None and mrow["published_month"] == month:
                last_val[mid] = (0, 0, 0)  # bu ay yayınlandı: ilk ölçümün tamamı kazanım
            else:
                last_val[mid] = cur  # ilk kez görülen eski gönderi: artış sayılmaz
        prev = last_val[mid]
        g = daily_gain[s["snapshot_date"]][group_of(mrow["content_type"] if mrow else None)]
        if cur[0] is not None:
            g.views_seen = True
        if cur[0] is not None and prev[0] is not None:
            g.views += max(0, cur[0] - prev[0])
        if cur[1] is not None and prev[1] is not None:
            g.likes += max(0, cur[1] - prev[1])
        if cur[2] is not None and prev[2] is not None:
            g.comments += max(0, cur[2] - prev[2])
        last_val[mid] = tuple(cur[k] if cur[k] is not None else prev[k] for k in range(3))
        latest[mid] = s

    # Bu ay yayınlanan gönderiler (son ölçülen değerleriyle)
    by_type = {t: TypeStats() for t in CONTENT_TYPES}
    by_group = {g: TypeStats() for g in GROUPS}
    posts: list[PostRow] = []
    for r in media_rows:
        if r["published_month"] != month:
            continue
        s = latest.get(r["media_id"])
        views = s["view_count"] if s else None
        likes = s["like_count"] if s else None
        comments = s["comments_count"] if s else None
        ct = r["content_type"] or "Diğer"
        by_type.setdefault(ct, TypeStats()).add_post(views, likes, comments)
        by_group[group_of(ct)].add_post(views, likes, comments)
        posts.append(PostRow(
            r["media_id"], ct, r["published_at"] or "", views, likes, comments,
            (r["caption"] or "").replace("\n", " ").strip(), r["permalink"] or "",
        ))
    posts.sort(key=lambda p: p.published_at)

    # Günlük tablo (ölçüm günleri ∪ paylaşım günleri)
    posts_by_day: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for p in posts:
        posts_by_day[p.published_at[:10]][p.content_type] += 1
    profile_by_day = {r["snapshot_date"]: r for r in series}
    days = sorted(set(profile_by_day) | set(posts_by_day) | set(daily_gain))
    prev_followers = baseline_profile["followers_count"] if baseline_profile else None
    daily: list[DailyRow] = []
    for day in days:
        prow = profile_by_day.get(day)
        followers = prow["followers_count"] if prow else None
        delta = followers - prev_followers if followers is not None and prev_followers is not None else None
        by_g = daily_gain.get(day) or {g: Gain() for g in GROUPS}
        total = Gain()
        for g in by_g.values():
            total.add(g)
        daily.append(DailyRow(day, followers, delta, dict(posts_by_day.get(day, {})), total, by_g))
        if followers is not None:
            prev_followers = followers

    gain = Gain()
    gain_by_group = {g: Gain() for g in GROUPS}
    for d in daily:
        gain.add(d.gain)
        for g in GROUPS:
            gain_by_group[g].add(d.gain_by_group[g])

    report = AccountReport(
        username=username,
        name=info["name"] if info else None,
        month=month,
        followers_start=start_row["followers_count"],
        followers_start_date=start_row["snapshot_date"],
        followers_end=end_row["followers_count"],
        followers_end_date=end_row["snapshot_date"],
        by_type=by_type,
        by_group=by_group,
        posts=posts,
        gain=gain,
        gain_by_group=gain_by_group,
        daily=daily,
        snapshot_days=len(series),
        view_count_available=any_view or not snaps,
    )
    if with_prev:
        p = build_account_report(conn, username, prev_month(month), with_prev=False)
        if p:
            report.prev = {"month": p.month, "posts": p.total_posts, "views": p.total_views,
                           "followers_delta": p.followers_delta, "gain": p.gain,
                           "gain_by_group": p.gain_by_group}
    return report


def ordered_usernames(conn, start: str, end: str) -> list[str]:
    """accounts.txt sırası + o dönemde verisi olan diğer hesaplar."""
    ordered = config.load_accounts()
    for u in db.usernames_with_data(conn, start, end):
        if u not in ordered:
            ordered.append(u)
    return ordered


def build_month_report(conn, month: str) -> list[AccountReport]:
    start, end = month_bounds(month)
    reports = []
    for username in ordered_usernames(conn, start, end):
        r = build_account_report(conn, username, month)
        if r:
            reports.append(r)
    return reports


# --- Markdown -------------------------------------------------------------------

def _split(total: str, reels: str, feed: str) -> str:
    return f"{total} ({reels} / {feed})"


def write_markdown(month: str, reports: list[AccountReport], path: Path) -> None:
    now = datetime.now()
    starts = [r.daily[0].day for r in reports if r.daily]
    ends = [r.daily[-1].day for r in reports if r.daily]
    lines = [
        f"# Instagram Rakip Raporu — {month_label(month)}",
        "",
        f"Oluşturulma: {now:%d.%m.%Y %H:%M} · Veri aralığı: {fmt_day(min(starts)) if starts else '–'} → "
        f"{fmt_day(max(ends)) if ends else '–'} · {len(reports)} hesap",
        "",
        "> Parantez içleri **(Reels / Feed)** kırılımıdır; Feed = fotoğraf + carousel + feed videosu. "
        "**Kazanılan**: ay içinde takip edilen tüm gönderilerin (önceki aylardan kalanlar dahil) "
        "izlenme/beğeni/yorum artışı. **Ay paylaşımları**: bu ay yayınlanan gönderilerin son ölçülen değerleri. "
        "Takipçi başlangıcı, bir önceki ayın son ölçümüdür.",
        "",
        "## Özet",
        "",
        "| Hesap | Takipçi | Takipçi Δ | Paylaşım (Reels / Feed) | Kazanılan izlenme (Reels / Feed) | "
        "Kazanılan beğeni (Reels / Feed) | Kazanılan yorum | Ort. beğeni Reels | Ort. beğeni Feed |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for r in reports:
        gr, gf = r.gain_by_group["Reels"], r.gain_by_group["Feed"]
        lines.append(
            f"| @{r.username} | {fmt(r.followers_end)} | {fmt_delta(r.followers_delta)} ({fmt_pct(r.followers_pct)}) | "
            f"{_split(str(r.total_posts), str(r.group_count('Reels')), str(r.group_count('Feed')))} | "
            f"{_split(fmt_delta(r.gain.views), fmt_delta(gr.views_or_none), fmt_delta(gf.views_or_none))} | "
            f"{_split(fmt_delta(r.gain.likes), fmt_delta(gr.likes), fmt_delta(gf.likes))} | "
            f"{fmt_delta(r.gain.comments)} | {fmt(r.by_group['Reels'].avg_likes)} | {fmt(r.by_group['Feed'].avg_likes)} |"
        )

    for r in reports:
        gr, gf = r.gain_by_group["Reels"], r.gain_by_group["Feed"]
        title = f"@{r.username}" + (f" — {r.name}" if r.name else "")
        lines += ["", f"## {title}", ""]
        lines.append(
            f"- **Takipçi:** {fmt(r.followers_start)} → {fmt(r.followers_end)} "
            f"(**{fmt_delta(r.followers_delta)}**, {fmt_pct(r.followers_pct)}) · "
            f"{fmt_day(r.followers_start_date)} → {fmt_day(r.followers_end_date)} · {r.snapshot_days} günlük ölçüm"
        )
        feed_parts = [f"{t} {r.type_count(t)}" for t in CONTENT_TYPES if t != "Reels" and r.type_count(t)]
        lines.append(
            f"- **Paylaşım:** {r.total_posts} · Reels {r.group_count('Reels')} · Feed {r.group_count('Feed')}"
            + (f" ({' · '.join(feed_parts)})" if feed_parts else "")
        )
        lines.append(
            f"- **Ay paylaşımlarının izlenmesi:** {fmt(r.total_views)} "
            f"(Reels {fmt(r.by_group['Reels'].views_or_none)} / Feed {fmt(r.by_group['Feed'].views_or_none)}) · "
            f"gönderi başına ort. {fmt(r.avg_views)}"
        )
        lines.append(
            f"- **Ay içinde kazanılan:** izlenme {fmt_delta(r.gain.views)} "
            f"(Reels {fmt_delta(gr.views_or_none)} / Feed {fmt_delta(gf.views_or_none)}) · "
            f"beğeni {fmt_delta(r.gain.likes)} (Reels {fmt_delta(gr.likes)} / Feed {fmt_delta(gf.likes)}) · "
            f"yorum {fmt_delta(r.gain.comments)}"
        )
        if r.prev:
            pg = r.prev["gain"]
            lines.append(
                f"- **Önceki ay ({month_label(r.prev['month'])}):** {r.prev['posts']} paylaşım · "
                f"kazanılan izlenme {fmt_delta(pg.views)} · kazanılan beğeni {fmt_delta(pg.likes)} · "
                f"takipçi {fmt_delta(r.prev['followers_delta'])}"
            )
        if not r.view_count_available:
            lines.append("- ⚠️ Bu hesap için API izlenme sayısı döndürmedi.")

        lines += ["", "| Reels / Feed | Paylaşım | Toplam izlenme | Ort. izlenme | Toplam beğeni | Ort. beğeni | "
                      "Kazanılan izlenme | Kazanılan beğeni | Kazanılan yorum |",
                  "|---|---:|---:|---:|---:|---:|---:|---:|---:|"]
        for g in GROUPS:
            ts, gn = r.by_group[g], r.gain_by_group[g]
            lines.append(f"| {g} | {ts.count} | {fmt(ts.views_or_none)} | {fmt(ts.avg_views)} | {fmt(ts.likes)} | "
                         f"{fmt(ts.avg_likes)} | {fmt_delta(gn.views_or_none)} | {fmt_delta(gn.likes)} | {fmt_delta(gn.comments)} |")

        if r.total_posts:
            lines += ["", "| Tür | Adet | Toplam izlenme | Ort. izlenme | Ort. beğeni | Ort. yorum |",
                      "|---|---:|---:|---:|---:|---:|"]
            for t in CONTENT_TYPES:
                ts = r.by_type.get(t)
                if ts and ts.count:
                    lines.append(f"| {t} | {ts.count} | {fmt(ts.views_or_none)} | {fmt(ts.avg_views)} | "
                                 f"{fmt(ts.avg_likes)} | {fmt(ts.avg_comments)} |")
            lines += ["", f"**En çok izlenen {min(TOP_N, r.total_posts)} gönderi**", "",
                      "| # | Tür | Tarih | İzlenme | Beğeni | Yorum | Açıklama |",
                      "|---:|---|---|---:|---:|---:|---|"]
            for i, p in enumerate(r.top_posts, 1):
                cap = p.caption[:60] + ("…" if len(p.caption) > 60 else "")
                cap = cap.replace("|", "/")
                link = f"[{cap or 'gönderi'}]({p.permalink})" if p.permalink else cap
                lines.append(f"| {i} | {p.content_type} | {fmt_day(p.published_at)} | {fmt(p.views)} | "
                             f"{fmt(p.likes)} | {fmt(p.comments)} | {link} |")
        else:
            lines.append("")
            lines.append("Bu ay paylaşım tespit edilmedi.")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# --- Excel ------------------------------------------------------------------------

def style_sheet(ws, widths: list[int] | None = None) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for i in range(1, ws.max_column + 1):
        w = widths[i - 1] if widths and i <= len(widths) else 13
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"


def write_xlsx(month: str, reports: list[AccountReport], path: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Özet"
    ws.append(["Hesap", "Ad", "Takipçi başlangıç", "Başlangıç tarihi", "Takipçi son", "Son tarih",
               "Takipçi Δ", "Takipçi %",
               "Paylaşım", "Reels paylaşım", "Feed paylaşım", "Fotoğraf", "Carousel", "Video",
               "İzlenme (ay paylaşımları)", "Reels izlenme", "Feed izlenme",
               "Ort. beğeni Reels", "Ort. beğeni Feed", "Ort. izlenme Reels", "Ort. izlenme Feed",
               "Kazanılan izlenme", "Kazanılan izlenme (Reels)", "Kazanılan izlenme (Feed)",
               "Kazanılan beğeni", "Kazanılan beğeni (Reels)", "Kazanılan beğeni (Feed)",
               "Kazanılan yorum", "Ölçüm günü"])
    for r in reports:
        br, bf = r.by_group["Reels"], r.by_group["Feed"]
        gr, gf = r.gain_by_group["Reels"], r.gain_by_group["Feed"]
        ws.append([r.username, r.name, r.followers_start, r.followers_start_date, r.followers_end,
                   r.followers_end_date, r.followers_delta,
                   round(r.followers_pct, 2) if r.followers_pct is not None else None,
                   r.total_posts, br.count, bf.count,
                   r.type_count("Fotoğraf"), r.type_count("Carousel"), r.type_count("Video"),
                   r.total_views, br.views_or_none, bf.views_or_none,
                   _round(br.avg_likes), _round(bf.avg_likes), _round(br.avg_views), _round(bf.avg_views),
                   r.gain.views, gr.views_or_none, gf.views_or_none,
                   r.gain.likes, gr.likes, gf.likes,
                   r.gain.comments, r.snapshot_days])
    style_sheet(ws, [18, 22, 13, 12, 13, 12, 11, 10, 10, 10, 10, 9, 9, 8, 15, 12, 12,
                     12, 12, 12, 12, 14, 14, 14, 14, 14, 14, 12, 9])

    ws = wb.create_sheet("Reels-Feed")
    ws.append(["Hesap", "Grup", "Paylaşım", "Toplam izlenme", "Ort. izlenme", "Toplam beğeni", "Ort. beğeni",
               "Toplam yorum", "Ort. yorum", "Kazanılan izlenme", "Kazanılan beğeni", "Kazanılan yorum"])
    for r in reports:
        for g in GROUPS:
            ts, gn = r.by_group[g], r.gain_by_group[g]
            ws.append([r.username, g, ts.count, ts.views_or_none, _round(ts.avg_views), ts.likes, _round(ts.avg_likes),
                       ts.comments, _round(ts.avg_comments), gn.views_or_none, gn.likes, gn.comments])
    style_sheet(ws, [18, 8, 10, 14, 12, 13, 11, 12, 10, 15, 15, 15])

    ws = wb.create_sheet("Türler")
    ws.append(["Hesap", "Tür", "Adet", "Toplam izlenme", "Ort. izlenme", "Toplam beğeni", "Ort. beğeni",
               "Toplam yorum", "Ort. yorum"])
    for r in reports:
        for t in CONTENT_TYPES:
            ts = r.by_type.get(t)
            if ts and ts.count:
                ws.append([r.username, t, ts.count, ts.views_or_none, _round(ts.avg_views),
                           ts.likes, _round(ts.avg_likes), ts.comments, _round(ts.avg_comments)])
    style_sheet(ws, [18, 10, 8, 14, 12, 14, 12, 12, 10])

    ws = wb.create_sheet("Günlük")
    ws.append(daily_header())
    for r in reports:
        for d in r.daily:
            ws.append(daily_row(r.username, d))
    style_sheet(ws, daily_widths())

    ws = wb.create_sheet("İçerikler")
    ws.append(["Hesap", "Tarih", "Grup", "Tür", "İzlenme", "Beğeni", "Yorum", "Açıklama", "Link"])
    for r in reports:
        for p in r.posts:
            ws.append([r.username, p.published_at[:16].replace("T", " "), p.group, p.content_type,
                       p.views, p.likes, p.comments, p.caption[:200], p.permalink])
    style_sheet(ws, [18, 17, 8, 10, 12, 10, 10, 60, 45])

    wb.save(path)


def daily_header() -> list[str]:
    return ["Hesap", "Tarih", "Takipçi", "Takipçi Δ", "Paylaşım", "Reels", "Feed",
            "Fotoğraf", "Carousel", "Video",
            "Kazanılan izlenme", "Kazanılan izlenme (Reels)", "Kazanılan izlenme (Feed)",
            "Kazanılan beğeni", "Kazanılan beğeni (Reels)", "Kazanılan beğeni (Feed)", "Kazanılan yorum"]


def daily_widths() -> list[int]:
    return [18, 12, 12, 11, 10, 8, 8, 9, 9, 8, 14, 14, 14, 14, 14, 14, 13]


def daily_row(username: str, d: DailyRow) -> list:
    gr, gf = d.gain_by_group["Reels"], d.gain_by_group["Feed"]
    return [username, d.day, d.followers, d.followers_delta, sum(d.posts.values()),
            d.posts_in_group("Reels"), d.posts_in_group("Feed"),
            d.posts.get("Fotoğraf", 0), d.posts.get("Carousel", 0), d.posts.get("Video", 0),
            d.gain.views, gr.views_or_none, gf.views_or_none, d.gain.likes, gr.likes, gf.likes, d.gain.comments]


def write_csv(month: str, reports: list[AccountReport], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["hesap", "tarih", "grup", "tur", "izlenme", "begeni", "yorum", "aciklama", "link", "media_id"])
        for r in reports:
            for p in r.posts:
                w.writerow([r.username, p.published_at, p.group, p.content_type, p.views, p.likes, p.comments,
                            p.caption[:300], p.permalink, p.media_id])


def generate(month: str, conn=None) -> dict | None:
    """Ay raporunu reports/YYYY-MM/ altına yazar. Veri yoksa None döner."""
    own = conn is None
    conn = conn or db.connect()
    try:
        reports = build_month_report(conn, month)
    finally:
        if own:
            conn.close()
    if not reports:
        log.info("%s için veri yok, rapor üretilmedi", month)
        return None
    out_dir = config.REPORTS_DIR / month
    out_dir.mkdir(parents=True, exist_ok=True)
    md = out_dir / f"rapor-{month}.md"
    xlsx = out_dir / f"rapor-{month}.xlsx"
    csv_path = out_dir / f"icerikler-{month}.csv"
    write_markdown(month, reports, md)
    write_xlsx(month, reports, xlsx)
    write_csv(month, reports, csv_path)
    log.info("Rapor yazıldı: %s (%d hesap)", out_dir, len(reports))
    return {"month": month, "dir": out_dir, "md": md, "xlsx": xlsx, "csv": csv_path, "reports": reports}
